"""Teilstudie B: ROI/TCO-Modell (DEMI 1.5, 3.5 und Tabelle 14).

Modellrechnung, keine Messung: Bis auf die variablen Kosten je Ticket (und für die Schwelle RAG vs. LLM-only
die Kosten von LLM-only sowie die Latenzdifferenz) sind alle Eingangsgrößen Sekundärdaten oder Annahmen.
Sämtliche Ergebnisse sind Modellergebnisse und keine empirischen Messwerte.

Aufruf im Repo-Hauptverzeichnis, nach auswertung/strang1.py:
    python auswertung/teilstudie_b.py

Eingabe: auswertung/ergebnisse/strang1/kennzahlen.csv (Mittelwerte aus dem Lauf haupt-01)

Ausgabe in auswertung/ergebnisse/teilstudie_b/:
- roi_modell.xlsx     Tabellenkalkulation mit Formeln: Hinweise, Parameter, Szenarien, Sensitivität
- bericht.md          Ergebnistabellen (deutsches Zahlenformat)
- szenarien.csv       Kennzahlen je Szenario
- sensitivitaet.csv   One-at-a-time-Variation ±20 % je Szenario
- break_even.png      Nutzen und Kosten über dem Ticketvolumen je Szenario
- tornado.png         Sensitivität des ROI je Szenario
- tornado_mittel.png  dasselbe nur für das mittlere Szenario (kompakte Fassung)
- meta.json           Herkunft: Eingabedatei mit Hash, Parameter, Git-Stand, Bibliotheksversionen
"""

import argparse
import csv
import hashlib
import json
import math
import sys
from pathlib import Path

AUSWERTUNG = Path(__file__).resolve().parent
ROOT = AUSWERTUNG.parent
sys.path.insert(0, str(ROOT / "rag"))
from common import git_info, now_iso  # noqa: E402

SZ = ("konservativ", "mittel", "optimistisch")
FARBE_NUTZEN, FARBE_KOSTEN = "#2a78d6", "#eb6834"      # Kategorialpalette, Slots 1 und 2
FARBE_PLUS, FARBE_MINUS = "#2a78d6", "#eb6834"
TEXT, TEXT2, GRID = "#0b0b0b", "#52514e", "#e4e3df"

# Kürzel, Bezeichnung, Einheit, Typ, Werte (konservativ, mittel, optimistisch) oder None (Messwert), Quelle, Zahlenformat
PARAMETER = [
    ("V", "Ticketvolumen je Monat", "Tickets", "Annahme", (2000, 10000, 30000),
     "DEMI, Tabelle 14; für den Break-even variiert", "#,##0"),
    ("a", "Adoptionsrate", "Anteil der Tickets", "Annahme", (0.40, 0.60, 0.80),
     "DEMI, Tabelle 14; gestützt auf Adhärenzbefunde (Brynjolfsson et al., 2025)", "0%"),
    ("t", "Zeitersparnis je unterstütztem Ticket", "min", "Annahme", (0.5, 1.5, 3.0),
     "DEMI, Tabelle 14; abgeleitet aus Brynjolfsson et al. (2025) und Xu et al. (2024)", "0.0"),
    ("p", "Personalkostensatz", "€/h", "Sekundärdaten", (33.20, 33.20, 33.20),
     "Destatis, Pressemitteilung Nr. 148 vom 29.04.2026: Arbeitskosten je geleistete Stunde 2025, "
     "Abschnitt N „Sonstige wirtschaftliche Dienstleistungen“ (B-03)", "#,##0.00"),
    ("r", "Reduktion der Eskalationsquote", "Prozentpunkte", "Annahme", (0.0, 0.02, 0.05),
     "DEMI, Tabelle 14; gestützt auf Mgbeahuruike et al. (2025)", "0.0%"),
    ("e", "Zusatzzeit je Eskalation", "min", "Annahme", (10, 15, 20),
     "DEMI, Tabelle 14; keine belastbare Referenz verfügbar", "0"),
    ("c", "Variable Kosten je Ticket (RAG)", "€", "Messwert", None,
     "Strang 1, Lauf haupt-01: Mittelwert RAG (B-05)", "0.000000"),
    ("s", "Serverpreis Hetzner CX33 (4 vCPU, 8 GB RAM, 80 GB)", "€/Monat", "Sekundärdaten", (8.49, 8.49, 8.49),
     "Hetzner Docs, „Hetzner Price Adjustment 15 June 2026“, Preis ab 15.06.2026, netto, abgerufen am 20.10.2026 (B-04)",
     "0.00"),
    ("ip", "Primäre IPv4-Adresse", "€/Monat", "Sekundärdaten", (0.50, 0.50, 0.50),
     "Hetzner Docs, „IPv4 pricing“, netto, abgerufen am 20.10.2026 (B-04)", "0.00"),
    ("bk", "Backup-Aufschlag", "Anteil am Serverpreis", "Sekundärdaten", (0.20, 0.20, 0.20),
     "Hetzner Docs, Cloud Billing FAQ, abgerufen am 20.10.2026 (B-04)", "0%"),
    ("hk", "Aufwand KB-Pflege", "h/Monat", "Annahme", (10, 5, 2),
     "DEMI, Tabelle 14; konservativ = höchster Aufwand (B-02)", "0"),
    ("hm", "Aufwand Monitoring", "h/Monat", "Annahme", (4, 2, 1),
     "DEMI, Tabelle 14; konservativ = höchster Aufwand (B-02)", "0"),
    ("cl", "Variable Kosten je Ticket (LLM-only)", "€", "Messwert", None,
     "Strang 1, Lauf haupt-01: Mittelwert LLM-only; nur für die Schwelle RAG vs. LLM-only (B-09)", "0.000000"),
    ("dl", "Latenzdifferenz RAG minus LLM-only", "s", "Messwert", None,
     "Strang 1, Lauf haupt-01: Mittelwert der Differenzen je Ticket; nur für die Schwelle (B-09)", "0.00"),
]
P_ZEILE = {k: 5 + i for i, (k, *_rest) in enumerate(PARAMETER)}       # Zeile im Blatt „Parameter“
P_SPALTE = dict(zip(SZ, "EFG"))                                         # Spalte im Blatt „Parameter“
S_SPALTE = dict(zip(SZ, "CDE"))                                         # Spalte im Blatt „Szenarien“

# One-at-a-time-Variation (DEMI 1.5): H = Hostingkosten insgesamt
SENS = [("V", "Ticketvolumen"), ("a", "Adoptionsrate"), ("t", "Zeitersparnis je Ticket"),
        ("p", "Personalkostensatz"), ("r", "Reduktion der Eskalationsquote"), ("e", "Zusatzzeit je Eskalation"),
        ("c", "Variable Kosten je Ticket"), ("H", "Hostingkosten"), ("hk", "Aufwand KB-Pflege"),
        ("hm", "Aufwand Monitoring")]
AMORT_MONATE = (12, 24, 36)


def de(x, nd=1) -> str:
    """Zahl im deutschen Format (Tausenderpunkt, Dezimalkomma, typografisches Minus)."""
    s = f"{x:,.{nd}f}".replace(",", "§").replace(".", ",").replace("§", ".")
    return s.replace("-", "−")


def eur(x, nd=2) -> str:
    return f"{de(x, nd)} €"


def pct(x, nd=0) -> str:
    return f"{de(x * 100, nd)} %"


def pp(x) -> str:
    """Spannweite des ROI in Prozentpunkten; kleine Werte mit einer Nachkommastelle."""
    return de(x * 100, 1 if x * 100 < 10 else 0)


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def lies_messwerte(path: Path) -> dict:
    """Mittelwerte aus strang1.py: Kosten je Anfrage (€) RAG und LLM-only, Latenzdifferenz je Ticket (s)."""
    if not path.exists():
        sys.exit(f"FEHLER: {path} fehlt. Zuerst auswertung/strang1.py ausführen.")
    with path.open(encoding="utf-8") as f:
        tab = {(r["bedingung"], r["groesse"]): r for r in csv.DictReader(f)}
    benoetigt = {"c": ("rag", "kosten_eur"), "cl": ("llm", "kosten_eur"), "dl": ("differenz_rag_minus_llm", "latenz_s")}
    werte = {}
    for k, key in benoetigt.items():
        if key not in tab:
            sys.exit(f"FEHLER: Zeile {key} fehlt in {path}")
        werte[k] = float(tab[key]["mittel"])
    if not (werte["c"] > 0 and werte["cl"] > 0):
        sys.exit("FEHLER: Kosten je Anfrage müssen positiv sein")
    return werte


def rechne(w: dict, f: dict | None = None) -> dict:
    """Modell nach DEMI 3.5; f enthält optional Faktoren je Kürzel für die Sensitivitätsanalyse."""
    f = f or {}

    def g(k):
        return w[k] * f.get(k, 1.0)

    V, a, t, p, r, e, c, hk, hm = (g(k) for k in ("V", "a", "t", "p", "r", "e", "c", "hk", "hm"))
    hosting = (w["s"] * (1 + w["bk"]) + w["ip"]) * f.get("H", 1.0)
    n = V * a
    nutzen_zeit = n * t / 60 * p
    nutzen_esk = n * r * e / 60 * p
    nutzen = nutzen_zeit + nutzen_esk
    kb_pflege, monitoring = hk * p, hm * p
    fix = hosting + kb_pflege + monitoring
    var = n * c
    kosten = fix + var
    netto = nutzen - kosten
    db = (t + r * e) / 60 * p - c                  # Deckungsbeitrag je unterstütztem Ticket
    be = fix / (a * db) if db > 0 else math.inf    # Break-even-Ticketvolumen
    mehr = hosting + kb_pflege + n * (c - w["cl"])  # Mehrkosten RAG gegenüber LLM-only (B-09)
    dt_kosten = mehr / (n * p) * 3600
    return {
        "unterstuetzte_tickets": n, "nutzen_zeit": nutzen_zeit, "nutzen_esk": nutzen_esk, "nutzen": nutzen,
        "hosting": hosting, "kb_pflege": kb_pflege, "monitoring": monitoring, "fixkosten": fix,
        "variable_kosten": var, "kosten": kosten, "nettonutzen": netto, "roi": netto / kosten,
        "anteil_variabel": var / kosten, "deckungsbeitrag": db, "break_even": be,
        **{f"amort_{m}": netto * m for m in AMORT_MONATE},
        "mehrkosten_rag": mehr, "schwelle_kosten_s": dt_kosten, "latenzdifferenz_s": w["dl"],
        "schwelle_gesamt_s": dt_kosten + w["dl"],
    }


# --------------------------------------------------------------------------- Tabellenkalkulation

def ref(k: str, sz: str) -> str:
    return f"Parameter!${P_SPALTE[sz]}${P_ZEILE[k]}"


def roi_formel(sz: str, fac_key: str | None = None, fac: str = "") -> str:
    """ROI als geschlossene Formel; fac_key wird mit dem Faktor fac multipliziert (Sensitivität)."""
    def v(k):
        return f"({ref(k, sz)}*{fac})" if k == fac_key else ref(k, sz)
    hosting = f"({ref('s', sz)}*(1+{ref('bk', sz)})+{ref('ip', sz)})"
    if fac_key == "H":
        hosting = f"({hosting}*{fac})"
    n = f"{v('V')}*{v('a')}"
    nutzen = f"{n}*({v('t')}+{v('r')}*{v('e')})/60*{v('p')}"
    kosten = f"({hosting}+({v('hk')}+{v('hm')})*{v('p')}+{n}*{v('c')})"
    return f"=({nutzen}-{kosten})/{kosten}"


SZEN_ZEILEN = [  # Schlüssel, Bezeichnung, Einheit, Formel (Platzhalter {X} = Parameter, [n] = Zeile im Blatt), Format
    ("unterstuetzte_tickets", "Unterstützte Tickets", "Tickets/Monat", "={V}*{a}", "#,##0"),
    ("nutzen_zeit", "Nutzen aus Zeitersparnis", "€/Monat", "=[unterstuetzte_tickets]*{t}/60*{p}", "#,##0.00"),
    ("nutzen_esk", "Nutzen aus vermiedenen Eskalationen", "€/Monat", "=[unterstuetzte_tickets]*{r}*{e}/60*{p}",
     "#,##0.00"),
    ("nutzen", "Nutzen gesamt", "€/Monat", "=[nutzen_zeit]+[nutzen_esk]", "#,##0.00"),
    ("hosting", "Hostingkosten (Server, IPv4, Backup)", "€/Monat", "={s}*(1+{bk})+{ip}", "#,##0.00"),
    ("kb_pflege", "KB-Pflege", "€/Monat", "={hk}*{p}", "#,##0.00"),
    ("monitoring", "Monitoring", "€/Monat", "={hm}*{p}", "#,##0.00"),
    ("fixkosten", "Fixkosten", "€/Monat", "=[hosting]+[kb_pflege]+[monitoring]", "#,##0.00"),
    ("variable_kosten", "Variable Kosten", "€/Monat", "=[unterstuetzte_tickets]*{c}", "#,##0.00"),
    ("kosten", "Kosten gesamt (TCO je Monat)", "€/Monat", "=[fixkosten]+[variable_kosten]", "#,##0.00"),
    ("nettonutzen", "Nettonutzen", "€/Monat", "=[nutzen]-[kosten]", "#,##0.00"),
    ("roi", "Modellierter ROI", "Anteil", "=[nettonutzen]/[kosten]", "0%"),
    ("anteil_variabel", "Anteil der variablen Kosten an den Kosten", "Anteil", "=[variable_kosten]/[kosten]", "0.00%"),
    ("deckungsbeitrag", "Deckungsbeitrag je unterstütztem Ticket", "€/Ticket", "=({t}+{r}*{e})/60*{p}-{c}", "0.0000"),
    ("break_even", "Break-even-Ticketvolumen", "Tickets/Monat",
     '=IF([deckungsbeitrag]>0,[fixkosten]/({a}*[deckungsbeitrag]),"nicht erreichbar")', "#,##0"),
    ("be_erreicht", "Break-even beim angenommenen Volumen erreicht", "", '=IF({V}>=[break_even],"ja","nein")', "@"),
    ("be_kontrolle", "Kontrolle: Nettonutzen beim Break-even-Volumen (soll 0)", "€/Monat",
     "=IFERROR([break_even]*{a}*[deckungsbeitrag]-[fixkosten],\"–\")", "0.00"),
    None,
    ("amort_12", "Max. Einmalinvestition, amortisiert in 12 Monaten", "€",
     '=IF([nettonutzen]>0,[nettonutzen]*12,"keine Amortisation")', "#,##0"),
    ("amort_24", "Max. Einmalinvestition, amortisiert in 24 Monaten", "€",
     '=IF([nettonutzen]>0,[nettonutzen]*24,"keine Amortisation")', "#,##0"),
    ("amort_36", "Max. Einmalinvestition, amortisiert in 36 Monaten", "€",
     '=IF([nettonutzen]>0,[nettonutzen]*36,"keine Amortisation")', "#,##0"),
    None,
    ("mehrkosten_rag", "Mehrkosten RAG gegenüber LLM-only", "€/Monat",
     "=[hosting]+[kb_pflege]+[unterstuetzte_tickets]*({c}-{cl})", "#,##0.00"),
    ("schwelle_kosten_s", "Erforderliche Zeitersparnis zur Deckung der Mehrkosten", "s je unterstütztem Ticket",
     "=[mehrkosten_rag]/([unterstuetzte_tickets]*{p})*3600", "0.00"),
    ("latenzdifferenz_s", "zuzüglich Latenzdifferenz (Wartezeit)", "s je unterstütztem Ticket", "={dl}", "0.00"),
    ("schwelle_gesamt_s", "Erforderliche zusätzliche Zeitersparnis gegenüber LLM-only", "s je unterstütztem Ticket",
     "=[schwelle_kosten_s]+[latenzdifferenz_s]", "0.00"),
]
ABSCHNITTE = {17: "Amortisationssicht (Ergänzung, B-08)", 21: "Schwelle RAG gegenüber LLM-only (Ergänzung, B-09)"}


def schreibe_xlsx(path: Path, werte_sz: dict, variation: float, stand: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    F = "Arial"
    f_norm, f_bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    f_titel, f_input = Font(name=F, size=12, bold=True), Font(name=F, size=10, color="0000FF")
    f_grau = Font(name=F, size=9, color="52514E")
    kopf_fill = PatternFill("solid", fgColor="EEEEEC")
    linie = Border(bottom=Side(style="thin", color="999999"))
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    # Hinweise
    ws = wb.active
    ws.title = "Hinweise"
    zeilen = [
        ("ROI/TCO-Modell – Teilstudie B", f_titel),
        ("Alle Ergebnisse sind Modellergebnisse und keine empirischen Messwerte (DEMI 1.5).", f_bold),
        ("", f_norm),
        ("Blatt „Parameter“: Eingangsgrößen je Szenario mit Typ und Quelle. Blaue Zahlen sind Eingabewerte, "
         "schwarze Zahlen Formeln.", f_norm),
        ("Blatt „Szenarien“: monatlicher Nutzen, Kosten (TCO), Nettonutzen, ROI, Break-even, Amortisationssicht "
         "und Schwelle RAG gegenüber LLM-only.", f_norm),
        (f"Blatt „Sensitivität“: One-at-a-time-Variation jedes Parameters um ±{de(variation * 100, 0)} % "
         "(Zelle C3 änderbar).", f_norm),
        ("", f_norm),
        ("Break-even per Zielwertsuche nachprüfen: Daten > Was-wäre-wenn-Analyse > Zielwertsuche; "
         "Zielzelle Szenarien!D15 (Nettonutzen, mittel), Zielwert 0, veränderbare Zelle Parameter!F5 (Ticketvolumen). "
         "Das Ergebnis entspricht Szenarien!D19. Danach Parameter!F5 auf den Ausgangswert zurücksetzen.", f_norm),
        ("", f_norm),
        (f"Erzeugt mit auswertung/teilstudie_b.py am {stand['erzeugt']}; Git-Stand {stand['git']}.", f_grau),
        (f"Messwerte aus {stand['eingabe']} (SHA-256 {stand['hash']}).", f_grau),
    ]
    for i, (txt, font) in enumerate(zeilen, start=1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font, c.alignment = font, Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    # Parameter
    wp = wb.create_sheet("Parameter")
    wp["A1"], wp["A1"].font = "Parameter des ROI/TCO-Modells (DEMI, Tabelle 14)", f_titel
    wp["A2"], wp["A2"].font = "Blau: Eingabewert. Schwarz: Formel (Wert gilt in allen Szenarien).", f_grau
    kopf = ["Kürzel", "Parameter", "Einheit", "Typ", "konservativ", "mittel", "optimistisch", "Quelle bzw. Annahme"]
    for j, h in enumerate(kopf, start=1):
        c = wp.cell(row=4, column=j, value=h)
        c.font, c.fill, c.border, c.alignment = f_bold, kopf_fill, linie, wrap
    for k, name, einheit, typ, _vals, quelle, fmt in PARAMETER:
        r = P_ZEILE[k]
        for j, v in enumerate([k, name, einheit, typ], start=1):
            wp.cell(row=r, column=j, value=v).font = f_norm
        gleich = len({werte_sz[sz][k] for sz in SZ}) == 1
        for j, sz in enumerate(SZ, start=5):
            c = wp.cell(row=r, column=j)
            if gleich and sz != "konservativ":
                c.value, c.font = f"=$E${r}", f_norm
            else:
                c.value, c.font = werte_sz[sz][k], f_input
            c.number_format = fmt
        q = wp.cell(row=r, column=8, value=quelle)
        q.font, q.alignment = f_norm, wrap
    for col, wdt in zip("ABCDEFGH", (8, 44, 20, 14, 13, 13, 13, 90)):
        wp.column_dimensions[col].width = wdt
    wp.freeze_panes = "E5"

    # Szenarien
    wsz = wb.create_sheet("Szenarien")
    wsz["A1"], wsz["A1"].font = "Monatliche Modellrechnung je Szenario (Modellergebnisse, keine Messwerte)", f_titel
    wsz["A2"], wsz["A2"].font = "Zeitangaben werden von Minuten in Stunden umgerechnet (/60).", f_grau
    for j, h in enumerate(["Größe", "Einheit", "konservativ", "mittel", "optimistisch"], start=1):
        c = wsz.cell(row=4, column=j, value=h)
        c.font, c.fill, c.border = f_bold, kopf_fill, linie
    zeile_von, r = {}, 5
    plan = []
    for i, item in enumerate(SZEN_ZEILEN):
        if item is None:
            r += 1
            plan.append((r, None, ABSCHNITTE[i]))
            r += 1
            continue
        zeile_von[item[0]] = r
        plan.append((r, item, None))
        r += 1
    for r, item, titel in plan:
        if item is None:
            wsz.cell(row=r, column=1, value=titel).font = f_bold
            continue
        key, name, einheit, formel, fmt = item
        wsz.cell(row=r, column=1, value=name).font = f_bold if key in ("kosten", "nettonutzen", "roi", "break_even") else f_norm
        wsz.cell(row=r, column=2, value=einheit).font = f_norm
        for sz in SZ:
            fx = formel
            for k, *_ in sorted(PARAMETER, key=lambda p: -len(p[0])):
                fx = fx.replace("{" + k + "}", ref(k, sz))
            for k2, zr in zeile_von.items():
                fx = fx.replace(f"[{k2}]", f"{S_SPALTE[sz]}{zr}")
            c = wsz.cell(row=r, column="CDE".index(S_SPALTE[sz]) + 3, value=fx)
            c.font, c.number_format = f_norm, fmt
            c.alignment = Alignment(horizontal="right")
    for col, wdt in zip("ABCDE", (58, 26, 16, 16, 16)):
        wsz.column_dimensions[col].width = wdt
    wsz.freeze_panes = "C5"
    assert zeile_von["nettonutzen"] == 15 and zeile_von["break_even"] == 19, "Hinweistext Zielwertsuche anpassen"

    # Sensitivität
    wse = wb.create_sheet("Sensitivität")
    wse["A1"], wse["A1"].font = "Sensitivitätsanalyse des modellierten ROI (One-at-a-time, DEMI 1.5)", f_titel
    wse["A3"], wse["A3"].font = "Variationsbreite", f_norm
    wse["C3"], wse["C3"].font, wse["C3"].number_format = variation, f_input, "0%"
    wse["D3"], wse["D3"].font = "jeder Parameter einzeln, alle übrigen auf dem Szenariowert", f_grau
    r = 5
    for sz in SZ:
        wse.cell(row=r, column=1, value=f"Szenario {sz}").font = f_bold
        r += 1
        for j, h in enumerate(["Parameter", "ROI bei −Variation", "ROI Basis", "ROI bei +Variation",
                               "Spannweite", "Rang"], start=1):
            c = wse.cell(row=r, column=j, value=h)
            c.font, c.fill, c.border = f_bold, kopf_fill, linie
        r += 1
        start = r
        for k, name in SENS:
            wse.cell(row=r, column=1, value=name).font = f_norm
            for j, fx in ((2, roi_formel(sz, k, "(1-$C$3)")),
                          (3, f"=Szenarien!{S_SPALTE[sz]}{zeile_von['roi']}"),
                          (4, roi_formel(sz, k, "(1+$C$3)")),
                          (5, f"=ROUND(ABS(D{r}-B{r}),9)")):
                c = wse.cell(row=r, column=j, value=fx)
                c.font, c.number_format = f_norm, "0.0%"
            c = wse.cell(row=r, column=6, value=f"=RANK(E{r},$E${start}:$E${start + len(SENS) - 1})")
            c.font = f_norm
            r += 1
        r += 1
    for col, wdt in zip("ABCDEF", (34, 20, 14, 20, 14, 8)):
        wse.column_dimensions[col].width = wdt

    wb["Hinweise"].sheet_view.showGridLines = False
    wb.save(path)


# --------------------------------------------------------------------------- Abbildungen

def abbildungen(out: Path, werte_sz: dict, erg: dict, sens: dict) -> list:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import numpy as np
    except ImportError:
        print("Hinweis: matplotlib fehlt, Abbildungen übersprungen (pip install matplotlib)")
        return []
    plt.rcParams.update({"font.size": 8.5, "axes.edgecolor": TEXT2, "axes.labelcolor": TEXT2,
                         "xtick.color": TEXT2, "ytick.color": TEXT2, "text.color": TEXT})
    fmt_ganz = matplotlib.ticker.FuncFormatter(lambda v, _: de(v, 0))

    # Break-even
    fig, axes = plt.subplots(1, 3, figsize=(6.3, 2.75), dpi=300, sharey=True)
    x = np.logspace(1, 5, 400)
    for ax, sz in zip(axes, SZ):
        w, e = werte_sz[sz], erg[sz]
        nutzen = x * w["a"] * (w["t"] + w["r"] * w["e"]) / 60 * w["p"]
        kosten = e["fixkosten"] + x * w["a"] * w["c"]
        ax.plot(x, nutzen, color=FARBE_NUTZEN, lw=1.6, label="Nutzen")
        ax.plot(x, kosten, color=FARBE_KOSTEN, lw=1.6, label="Kosten (TCO)")
        ax.set_xscale("log")
        ax.set_yscale("log")
        ax.axvline(w["V"], color=TEXT2, lw=0.9, ls=(0, (3, 3)))
        ax.text(w["V"] * 1.12, 1.6, f"Szenario-\nvolumen\n{de(w['V'], 0)}", color=TEXT2, fontsize=7, va="bottom")
        be = e["break_even"]
        if math.isfinite(be):
            ax.plot([be], [e["fixkosten"] + be * w["a"] * w["c"]], "o", ms=5.5, color=TEXT, mec="white", mew=1.2,
                    zorder=5)
            ax.annotate(f"Break-even\n{de(math.ceil(be), 0)}", (be, e["fixkosten"] + be * w["a"] * w["c"]),
                        xytext=(5, -6), textcoords="offset points", ha="left", va="top", fontsize=7, color=TEXT)
        ax.set_title(sz, fontsize=9, color=TEXT, loc="left")
        ax.set_xlim(10, 1e5)
        ax.set_ylim(1, 3e5)
        ax.grid(True, which="major", color=GRID, linewidth=0.6)
        ax.set_axisbelow(True)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        ax.xaxis.set_major_formatter(fmt_ganz)
        ax.yaxis.set_major_formatter(fmt_ganz)
        ax.tick_params(axis="x", labelsize=7)
    axes[0].set_ylabel("€ je Monat (log. Skala)")
    axes[1].set_xlabel("Ticketvolumen je Monat (log. Skala)")
    h, lab = axes[0].get_legend_handles_labels()
    fig.legend(h, lab, loc="upper center", ncol=2, frameon=False, fontsize=7.5, bbox_to_anchor=(0.5, 1.0))
    fig.tight_layout(rect=(0, 0, 1, 0.93))
    fig.savefig(out / "break_even.png", facecolor="white")
    plt.close(fig)

    # Tornado
    fig, axes = plt.subplots(3, 1, figsize=(6.3, 7.4), dpi=300)
    for ax, sz in zip(axes, SZ):
        rows = sorted(sens[sz], key=lambda s: -s["spannweite"])
        basis = erg[sz]["roi"] * 100
        y = np.arange(len(rows))[::-1]
        for yi, s in zip(y, rows):
            lo, hi = s["roi_minus"] * 100, s["roi_plus"] * 100
            ax.barh(yi + 0.17, lo - basis, left=basis, height=0.32, color=FARBE_MINUS,
                    label="Parameter −20 %" if yi == y[0] else None)
            ax.barh(yi - 0.17, hi - basis, left=basis, height=0.32, color=FARBE_PLUS,
                    label="Parameter +20 %" if yi == y[0] else None)
        ax.axvline(basis, color=TEXT, lw=0.9)
        ax.set_yticks(y, [s["name"] for s in rows], fontsize=7.5)
        ax.set_title(f"{sz} (Basis-ROI {de(basis, 0)} %)", fontsize=9, color=TEXT, loc="left")
        ax.xaxis.grid(True, color=GRID, linewidth=0.6)
        ax.set_axisbelow(True)
        for side in ("top", "right", "left"):
            ax.spines[side].set_visible(False)
        ax.tick_params(axis="y", length=0)
        ax.xaxis.set_major_locator(matplotlib.ticker.MaxNLocator(nbins=6, steps=[1, 2, 5, 10]))
        ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{de(v, 0)} %"))
        lo_all = min(min(s["roi_minus"], s["roi_plus"]) for s in rows) * 100
        hi_all = max(max(s["roi_minus"], s["roi_plus"]) for s in rows) * 100
        pad = (hi_all - lo_all) * 0.06 or 1
        ax.set_xlim(lo_all - pad, hi_all + pad)
    axes[0].legend(loc="lower right", frameon=False, fontsize=7.5)
    axes[-1].set_xlabel("modellierter ROI")
    fig.tight_layout()
    fig.savefig(out / "tornado.png", facecolor="white")
    plt.close(fig)

    # Tornado nur für das mittlere Szenario (kompakte Fassung für das EVI)
    fig, ax = plt.subplots(figsize=(6.3, 2.9), dpi=300)
    rows = sorted(sens["mittel"], key=lambda s: -s["spannweite"])
    basis = erg["mittel"]["roi"] * 100
    y = np.arange(len(rows))[::-1]
    for yi, s in zip(y, rows):
        lo, hi = s["roi_minus"] * 100, s["roi_plus"] * 100
        ax.barh(yi + 0.17, lo - basis, left=basis, height=0.32, color=FARBE_MINUS,
                label="Parameter −20 %" if yi == y[0] else None)
        ax.barh(yi - 0.17, hi - basis, left=basis, height=0.32, color=FARBE_PLUS,
                label="Parameter +20 %" if yi == y[0] else None)
    ax.axvline(basis, color=TEXT, lw=0.9)
    ax.set_yticks(y, [s["name"] for s in rows], fontsize=7.5)
    ax.xaxis.grid(True, color=GRID, linewidth=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.xaxis.set_major_locator(matplotlib.ticker.MaxNLocator(nbins=6, steps=[1, 2, 5, 10]))
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{de(v, 0)} %"))
    ax.set_xlabel(f"modellierter ROI im mittleren Szenario (Basis {de(basis, 0)} %)")
    ax.legend(loc="lower right", frameon=False, fontsize=7.5)
    fig.tight_layout()
    fig.savefig(out / "tornado_mittel.png", facecolor="white")
    plt.close(fig)
    return ["break_even.png", "tornado.png", "tornado_mittel.png"]


# --------------------------------------------------------------------------- Hauptprogramm

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--kennzahlen", type=Path, default=AUSWERTUNG / "ergebnisse" / "strang1" / "kennzahlen.csv")
    ap.add_argument("--out", type=Path, default=AUSWERTUNG / "ergebnisse" / "teilstudie_b")
    ap.add_argument("--variation", type=float, default=0.20, help="Variationsbreite der Sensitivitätsanalyse")
    args = ap.parse_args()

    mess = lies_messwerte(args.kennzahlen)
    werte_sz = {}
    for i, sz in enumerate(SZ):
        werte_sz[sz] = {k: (mess[k] if vals is None else vals[i]) for k, _n, _e, _t, vals, _q, _f in PARAMETER}
    erg = {sz: rechne(werte_sz[sz]) for sz in SZ}

    sens = {}
    for sz in SZ:
        sens[sz] = []
        for k, name in SENS:
            lo = rechne(werte_sz[sz], {k: 1 - args.variation})["roi"]
            hi = rechne(werte_sz[sz], {k: 1 + args.variation})["roi"]
            sens[sz].append({"kuerzel": k, "name": name, "roi_minus": lo, "roi_basis": erg[sz]["roi"],
                             "roi_plus": hi, "spannweite": abs(hi - lo)})
        for s in sens[sz]:   # Wettkampfrang wie RANK() in der Tabellenkalkulation: gleiche Spannweite, gleicher Rang
            s["rang"] = 1 + sum(round(o["spannweite"], 9) > round(s["spannweite"], 9) for o in sens[sz])

    out = args.out
    out.mkdir(parents=True, exist_ok=True)
    git_state = git_info()
    try:
        eingabe = str(args.kennzahlen.resolve().relative_to(ROOT))
    except ValueError:
        eingabe = str(args.kennzahlen)
    stand = {"erzeugt": now_iso(), "git": (git_state.get("commit") or "unbekannt")[:12]
             + (" (nicht committete Änderungen)" if git_state.get("dirty") else ""),
             "eingabe": eingabe, "hash": file_hash(args.kennzahlen)}

    schreibe_xlsx(out / "roi_modell.xlsx", werte_sz, args.variation, stand)

    with (out / "szenarien.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["groesse"] + list(SZ))
        for key in erg["mittel"]:
            w.writerow([key] + [round(erg[sz][key], 8) if math.isfinite(erg[sz][key]) else "" for sz in SZ])
    with (out / "sensitivitaet.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["szenario", "kuerzel", "parameter", "roi_minus", "roi_basis", "roi_plus", "spannweite", "rang"])
        for sz in SZ:
            for s in sorted(sens[sz], key=lambda s: s["rang"]):
                w.writerow([sz, s["kuerzel"], s["name"]] + [round(s[k], 6) for k in
                                                            ("roi_minus", "roi_basis", "roi_plus", "spannweite")]
                           + [s["rang"]])

    bilder = abbildungen(out, werte_sz, erg, sens)

    # Bericht
    L = [
        "# Teilstudie B: ROI/TCO-Modell",
        "",
        f"Erzeugt mit `auswertung/teilstudie_b.py` am {stand['erzeugt'][:10]}. **Alle Werte sind Modellergebnisse "
        "und keine empirischen Messwerte** (DEMI 1.5). Gemessen sind nur die variablen Kosten je Ticket und für die "
        "Schwelle RAG gegenüber LLM-only die Latenzdifferenz (Strang 1); alle übrigen Eingangsgrößen sind "
        "Sekundärdaten oder Annahmen. Beträge netto, je Monat.",
        "",
        "## Eingangsgrößen",
        "",
        "| Parameter | Einheit | Typ | konservativ | mittel | optimistisch |",
        "| --- | --- | --- | --- | --- | --- |",
    ]

    def pfmt(k, v):
        if k in ("a", "bk"):
            return pct(v, 0)
        if k == "r":
            return de(v * 100, 0)
        if k in ("c", "cl"):
            return de(v, 6)
        if k in ("s", "ip", "p", "dl"):
            return de(v, 2)
        if k == "t":
            return de(v, 1)
        return de(v, 0)

    for k, name, einheit, typ, _v, _q, _f in PARAMETER:
        L.append(f"| {name} | {einheit} | {typ} | " + " | ".join(pfmt(k, werte_sz[sz][k]) for sz in SZ) + " |")
    L += ["", "Quellen und Begründungen: `auswertung/README.md`, Entscheidungen B-01 bis B-09.", "",
          "## Tabelle: Monatlicher Nutzen, Kosten und ROI je Szenario", "",
          "| Größe | konservativ | mittel | optimistisch |", "| --- | --- | --- | --- |"]

    def zeile(label, key, f=eur, fett=False):
        txt = " | ".join(f(erg[sz][key]) for sz in SZ)
        L.append(f"| **{label}** | {txt} |" if fett else f"| {label} | {txt} |")

    zeile("Unterstützte Tickets", "unterstuetzte_tickets", lambda v: de(v, 0))
    zeile("Nutzen aus Zeitersparnis", "nutzen_zeit")
    zeile("Nutzen aus vermiedenen Eskalationen", "nutzen_esk")
    zeile("Nutzen gesamt", "nutzen", fett=True)
    zeile("Hostingkosten", "hosting")
    zeile("KB-Pflege", "kb_pflege")
    zeile("Monitoring", "monitoring")
    zeile("Variable Kosten (Tokens, Embeddings)", "variable_kosten")
    zeile("Kosten gesamt (TCO)", "kosten", fett=True)
    zeile("Nettonutzen", "nettonutzen", fett=True)
    zeile("Modellierter ROI", "roi", pct, fett=True)
    zeile("Anteil variable Kosten an den Kosten", "anteil_variabel", lambda v: pct(v, 2))
    L += ["", "## Break-even", "",
          "Ticketvolumen je Monat, ab dem der Nutzen die Kosten deckt (übrige Parameter auf dem Szenariowert; "
          "aufgerundet auf ganze Tickets). Das Modell ist im Ticketvolumen linear; der Wert wird daher geschlossen "
          "berechnet und entspricht dem Ergebnis der Zielwertsuche (B-06).", "",
          "| | konservativ | mittel | optimistisch |", "| --- | --- | --- | --- |",
          "| Break-even-Ticketvolumen | " + " | ".join(
              de(math.ceil(erg[sz]["break_even"]), 0) if math.isfinite(erg[sz]["break_even"]) else "nicht erreichbar"
              for sz in SZ) + " |",
          "| angenommenes Ticketvolumen | " + " | ".join(de(werte_sz[sz]["V"], 0) for sz in SZ) + " |",
          "| Break-even erreicht | " + " | ".join("ja" if werte_sz[sz]["V"] >= erg[sz]["break_even"] else "nein"
                                                 for sz in SZ) + " |",
          "", "## Amortisationssicht (Ergänzung, B-08)", "",
          "Höchste einmalige Investition (Einführung, Integration, Schulung), die der monatliche Nettonutzen in der "
          "jeweiligen Frist zurückverdient; ohne Abzinsung.", "",
          "| Frist | konservativ | mittel | optimistisch |", "| --- | --- | --- | --- |"]
    for mo in AMORT_MONATE:
        L.append(f"| {mo} Monate | " + " | ".join(
            eur(erg[sz][f"amort_{mo}"], 0) if erg[sz]["nettonutzen"] > 0 else "keine Amortisation" for sz in SZ) + " |")
    L += ["", "## Schwelle RAG gegenüber LLM-only (Ergänzung, B-09)", "",
          "Zusätzliche Zeitersparnis je unterstütztem Ticket, die RAG gegenüber LLM-only erbringen muss, damit sich "
          "Hosting, KB-Pflege und die höheren Tokenkosten tragen; zuzüglich der gemessenen Latenzdifferenz als "
          "Wartezeit.", "",
          "| Größe | konservativ | mittel | optimistisch |", "| --- | --- | --- | --- |"]
    zeile("Mehrkosten RAG je Monat", "mehrkosten_rag")
    zeile("Zeitersparnis zur Deckung der Mehrkosten (s/Ticket)", "schwelle_kosten_s", lambda v: de(v, 1))
    zeile("Latenzdifferenz (s/Ticket)", "latenzdifferenz_s", lambda v: de(v, 2))
    zeile("Erforderliche zusätzliche Zeitersparnis (s/Ticket)", "schwelle_gesamt_s", lambda v: de(v, 1), fett=True)
    L += ["", f"## Sensitivität (One-at-a-time, ±{de(args.variation * 100, 0)} %)", "",
          "Rangfolge nach der Spannweite des ROI zwischen −20 % und +20 %; alle übrigen Parameter auf dem "
          "Szenariowert.", "",
          "| konservativ | mittel | optimistisch |", "| --- | --- | --- |"]
    geordnet = {sz: sorted(sens[sz], key=lambda s: (s["rang"], SENS.index((s["kuerzel"], s["name"])))) for sz in SZ}
    for i in range(len(SENS)):
        L.append("| " + " | ".join(
            f"{geordnet[sz][i]['rang']}. {geordnet[sz][i]['name']} ({pp(geordnet[sz][i]['spannweite'])} Pp.)"
            for sz in SZ) + " |")
    L += ["", "Pp. = Prozentpunkte des ROI; gleiche Spannweite ergibt denselben Rang. Ticketvolumen und Adoptionsrate "
          "gehen nur als Produkt (unterstützte Tickets) in das Modell ein und wirken deshalb gleich stark; dasselbe gilt "
          "für Reduktion der Eskalationsquote und Zusatzzeit je Eskalation. Im konservativen Szenario ist die Reduktion der "
          "Eskalationsquote 0; beide Eskalationsparameter haben dort bei relativer Variation keine Wirkung.", "",
          "Szenario mittel im Einzelnen:", "",
          "| Parameter | ROI bei −20 % | ROI bei +20 % | Spannweite (Pp.) |", "| --- | --- | --- | --- |"]
    for s in geordnet["mittel"]:
        L.append(f"| {s['name']} | {pct(s['roi_minus'])} | {pct(s['roi_plus'])} | {pp(s['spannweite'])} |")
    L.append("")
    for b in bilder:
        L += [f"![{b}]({b})", ""]
    (out / "bericht.md").write_text("\n".join(L), encoding="utf-8")

    versionen = {"python": sys.version.split()[0]}
    for mod in ("openpyxl", "matplotlib", "numpy"):
        try:
            versionen[mod] = __import__(mod).__version__
        except ImportError:
            pass
    s1_meta = args.kennzahlen.parent / "meta.json"
    meta = {"erzeugt": stand["erzeugt"], "skript": "auswertung/teilstudie_b.py", "git": git_state,
            "eingabe": {"datei": eingabe, "sha256": stand["hash"],
                        "strang1_meta": json.loads(s1_meta.read_text(encoding="utf-8")) if s1_meta.exists() else None},
            "variation": args.variation, "parameter": werte_sz, "versionen": versionen}
    (out / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print("\n".join(L[:3]))
    print("ROI: " + ", ".join(f"{sz} {pct(erg[sz]['roi'])}" for sz in SZ))
    print("Break-even: " + ", ".join(f"{sz} {de(math.ceil(erg[sz]['break_even']), 0)}" for sz in SZ))
    try:
        rel = out.resolve().relative_to(ROOT)
    except ValueError:
        rel = out
    print(f"\nGeschrieben: {rel}/ (roi_modell.xlsx, bericht.md, szenarien.csv, sensitivitaet.csv, "
          f"{', '.join(bilder)}, meta.json)")


if __name__ == "__main__":
    main()
