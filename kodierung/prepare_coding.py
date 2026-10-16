"""Erzeugt Kodierheft und Kodiertabelle für die manuelle Kodierung (Strang 4, DEMI 2.5 und 3.3; Leitfaden: LEITFADEN.md).

Eingabe:
- rag/runs/<run-id>/responses.jsonl   Antwortentwürfe mit übergebenen Auszügen (Standard: haupt-01)
- tickets.jsonl                       Musterantwort, Kernaussagen, Belege und Belegartikel je Ticket (Hilfsmittel)
- kb/KB-*.md                          vollständige Knowledge Base (Maßstab für C1)

Ausgabe (Standard: kodierung/, anderer Ordner mit --out):
- kodierheft.html          alles, was zum Kodieren nötig ist, in Kodierreihenfolge; KB-Artikel im Anhang verlinkt
- kodiertabelle.xlsx       eine Zeile je Entwurf mit Auswahllisten für die Codes
- stichprobe_zweitkodierung.csv   48 Entwürfe (24 RAG, 24 LLM-only) für die Zweitkodierung (nur Hauptlauf, K-06)
- kodier_meta.json         Herkunft: Lauf, Hashes, Seed, Git-Stand

Aufruf (im Repo-Hauptverzeichnis; benötigt openpyxl):
    python kodierung/prepare_coding.py                  # Hauptlauf haupt-01, 240 Entwürfe
"""

import argparse
import csv
import hashlib
import html
import json
import random
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("FEHLER: openpyxl fehlt. Installation: pip install openpyxl")

KODIER_DIR = Path(__file__).resolve().parent
ROOT = KODIER_DIR.parent
sys.path.insert(0, str(ROOT / "rag"))
from common import git_info, now_iso  # noqa: E402

SEED = 20261015                  # Kodierreihenfolge und Stichprobe der Zweitkodierung (K-05, K-06)
N_ZWEIT = {"rag": 24, "llm": 24}  # 20 % von 240, je Bedingung gleich viele (K-06)
CIT = re.compile(r"\[\s*KB-\d{3}(?:\s*[,;]\s*KB-\d{3})*\s*\]")
KB_ID = re.compile(r"KB-\d{3}")

C1 = ["gedeckt", "teilweise", "nicht gedeckt"]
C2 = ["passend", "nicht passend", "kein Beleg"]
JA_NEIN = ["ja", "nein"]
FEHLER = ["F1", "F2", "F3", "F4", "keiner"]
NA = "n. a."


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def esc(text: str) -> str:
    return html.escape(text or "", quote=True)


def link_kb(match: re.Match) -> str:
    return KB_ID.sub(lambda m: f'<a href="#{m.group(0)}">{m.group(0)}</a>', match.group(0))


def answer_html(text: str) -> str:
    """Belege hervorheben und auf den KB-Artikel im Anhang verlinken."""
    escaped = esc(text)
    return CIT.sub(lambda m: f"<mark>{link_kb(m)}</mark>", escaped)


def kb_links(ids) -> str:
    return ", ".join(f'<a href="#{i}">{i}</a>' for i in ids) or "–"


def load_kb() -> dict:
    kb = {}
    for path in sorted((ROOT / "kb").glob("KB-*.md")):
        text = path.read_text(encoding="utf-8")
        body = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S).strip()
        title = re.search(r"^# (.+)$", body, flags=re.M)
        kb[path.stem] = {"titel": title.group(1) if title else path.stem, "text": body}
    if len(kb) != 40:
        sys.exit(f"FEHLER: {len(kb)} KB-Artikel gefunden, erwartet 40.")
    return kb


def build_html(items: list, tickets: dict, kb: dict, run_id: str) -> str:
    css = """
    body{font-family:system-ui,-apple-system,Segoe UI,sans-serif;max-width:980px;margin:0 auto;padding:16px;line-height:1.5;color:#1d1d1f;background:#fff}
    h1{font-size:1.5rem} h2{font-size:1.15rem;margin:0 0 8px} h3{font-size:.95rem;margin:16px 0 4px;color:#444}
    section.d{border-top:3px solid #333;padding:16px 0 24px;margin-top:24px}
    .txt{white-space:pre-wrap;background:#f6f6f6;border:1px solid #ddd;border-radius:6px;padding:10px 12px}
    .ans{background:#fffdf3;border-color:#e6dca8}
    .pas{white-space:pre-wrap;background:#f2f6fb;border:1px solid #c9d6e6;border-radius:6px;padding:8px 12px;margin:6px 0;font-size:.92rem}
    .meta{color:#555;font-size:.9rem} .ok{color:#137333;font-weight:600} .no{color:#b3261e;font-weight:600}
    mark{background:#ffe58a;padding:0 2px} table{border-collapse:collapse;width:100%;font-size:.9rem}
    td,th{border:1px solid #ddd;padding:4px 6px;text-align:left} nav{margin:12px 0;font-size:.9rem}
    details{margin-top:6px} summary{cursor:pointer;color:#0b57d0}
    .kb{white-space:pre-wrap;font-size:.9rem;background:#fafafa;border:1px solid #ddd;border-radius:6px;padding:10px 12px}
    .tag{display:inline-block;padding:1px 8px;border-radius:10px;font-size:.8rem;font-weight:600;margin-left:6px}
    .rag{background:#dfe9fb;color:#0b3d91} .llm{background:#eee;color:#333} .gap{background:#fde2e1;color:#8c1d18}
    """
    out = [f"<!doctype html><html lang='de'><head><meta charset='utf-8'>"
           f"<meta name='viewport' content='width=device-width,initial-scale=1'><title>Kodierheft {esc(run_id)}</title>"
           f"<style>{css}</style></head><body>",
           f"<h1 id='top'>Kodierheft – Lauf {esc(run_id)}</h1>",
           "<p class='meta'>Reihenfolge = Kodierreihenfolge (K-05). Maßstab für C1 ist die gesamte Knowledge Base im Anhang, "
           "nicht die Musterantwort. Regeln: LEITFADEN.md. Codes in kodiertabelle.xlsx eintragen.</p>",
           "<table><tr><th>Kodier-ID</th><th>Ticket</th><th>Bedingung</th><th>KB-Abdeckung</th></tr>"]
    for it in items:
        out.append(f"<tr><td><a href='#{it['kodier_id']}'>{it['kodier_id']}</a></td><td>{esc(it['ticket_id'])}</td>"
                   f"<td>{'RAG' if it['bedingung'] == 'rag' else 'LLM-only'}</td><td>{esc(it['kb_abdeckung'])}</td></tr>")
    out.append("</table><p><a href='#kb'>→ Knowledge Base (Anhang)</a></p>")

    for n, it in enumerate(items):
        e, t = it["entry"], tickets[it["ticket_id"]]
        prev_id = items[n - 1]["kodier_id"] if n else None
        next_id = items[n + 1]["kodier_id"] if n + 1 < len(items) else None
        bed = "rag" if it["bedingung"] == "rag" else "llm"
        tags = f"<span class='tag {bed}'>{'RAG' if bed == 'rag' else 'LLM-only'}</span>"
        if e["kb_abdeckung"] == "gap":
            tags += "<span class='tag gap'>KB-Gap</span>"
        out.append(f"<section class='d' id='{it['kodier_id']}'>")
        out.append(f"<h2>{it['kodier_id']} · {esc(e['ticket_id'])}{tags}</h2>")
        out.append(f"<div class='meta'>{esc(e['kategorie'])} · {esc(e['schwierigkeit'])} · KB-Abdeckung: {esc(e['kb_abdeckung'])}"
                   + (f" · Gap-Thema: {esc(t.get('gap_thema', ''))}" if e["kb_abdeckung"] == "gap" else "") + "</div>")
        out.append(f"<h3>Kundenanfrage</h3><div class='txt'>{esc(e['ticket_text'])}</div>")
        out.append(f"<h3>Antwortentwurf</h3><div class='txt ans'>{answer_html(e['antwort'])}</div>")

        if bed == "rag":
            passages = (e.get("retrieval") or {}).get("passagen", [])
            passed = {p["artikel_id"] for p in passages}
            cands = (e.get("retrieval") or {}).get("kandidaten", [])
            out.append("<h3>Prüfhilfe Quellen</h3><table>")
            out.append(f"<tr><td>Zitierte Artikel</td><td>{kb_links(e.get('zitierte_artikel', []))}</td></tr>")
            outside = e.get("zitate_ausserhalb_kontext", [])
            out.append(f"<tr><td>Zitiert, aber nicht übergeben</td><td>{kb_links(outside) if outside else '–'}</td></tr>")
            marks = " · ".join(f"{a} <span class='{'ok' if a in passed else 'no'}'>{'übergeben' if a in passed else 'nicht übergeben'}</span>"
                               for a in t.get("belegartikel", [])) or "– (keine Belegartikel)"
            out.append(f"<tr><td>Belegartikel laut Ticket</td><td>{marks}</td></tr>")
            weitere = t.get("weitere_kandidaten", [])
            if weitere:
                out.append(f"<tr><td>Weitere relevante Artikel</td><td>{kb_links(weitere)}</td></tr>")
            out.append(f"<tr><td>Abgerufene Kandidaten (Top-5)</td><td>{kb_links([c['chunk_id'].split('#')[0] for c in cands])} "
                       f"<span class='meta'>({', '.join(esc(c['chunk_id']) for c in cands)})</span></td></tr></table>")
            out.append("<h3>Übergebene Auszüge (Top-3 nach Reranking)</h3>")
            for p in passages:
                out.append(f"<div class='pas'><b>{p.get('rank_rerank', '')}. <a href='#{p['artikel_id']}'>{p['artikel_id']}</a> "
                           f"– {esc(p.get('titel', ''))}</b> <span class='meta'>({esc(p['chunk_id'])})</span>\n{esc(p['text'])}</div>")
            if not passages:
                out.append("<p class='no'>Keine Auszüge im Protokoll.</p>")

        out.append("<h3>Referenz (Hilfsmittel, nicht Maßstab für C1)</h3><table>")
        out.append(f"<tr><td>Antworttyp</td><td>{esc(t.get('antworttyp', ''))}</td></tr>")
        ka = t.get("kernaussagen") or []
        out.append("<tr><td>Kernaussagen der Musterantwort</td><td>"
                   + ("<br>".join(f"• {esc(k)}" for k in ka) if ka else "– (keine; Rückfrage bzw. Enthaltung erwartet)") + "</td></tr>")
        bel = t.get("belege") or []
        if bel:
            out.append("<tr><td>Belegzitate</td><td>" + "<br>".join(
                f"<a href='#{b['artikel_id']}'>{b['artikel_id']}</a>: „{esc(b['zitat'])}“" for b in bel) + "</td></tr>")
        out.append("</table>")
        out.append(f"<details><summary>Musterantwort anzeigen</summary><div class='txt'>{esc(t.get('musterantwort', ''))}</div></details>")
        nav = [f"<a href='#{prev_id}'>← {prev_id}</a>" if prev_id else "", "<a href='#top'>Übersicht</a>",
               f"<a href='#{next_id}'>{next_id} →</a>" if next_id else ""]
        out.append(f"<nav>{' · '.join(x for x in nav if x)}</nav></section>")

    out.append("<section id='kb'><h1>Anhang: Knowledge Base (kb-v1.0)</h1>")
    for kid, a in kb.items():
        out.append(f"<section id='{kid}'><h2>{kid} – {esc(a['titel'])}</h2><div class='kb'>{esc(a['text'])}</div>"
                   f"<nav><a href='#top'>Übersicht</a></nav></section>")
    out.append("</section></body></html>")
    return "\n".join(out)


def build_xlsx(items: list, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kodierung"
    cols = ["kodier_id", "ticket_id", "bedingung", "kb_abdeckung", "schwierigkeit",
            "K1", "K1_C1", "K1_C2", "K2", "K2_C1", "K2_C2", "K3", "K3_C1", "K3_C2",
            "C3", "mangel", "F_primaer", "F_weitere", "bemerkung"]
    ws.append(cols)
    grey = PatternFill("solid", fgColor="E7E7E7")
    for c in ws[1]:
        c.font = Font(bold=True)
    col = {name: i + 1 for i, name in enumerate(cols)}
    for r, it in enumerate(items, start=2):
        llm = it["bedingung"] == "llm"
        ws.cell(r, col["kodier_id"], it["kodier_id"]).hyperlink = f"kodierheft.html#{it['kodier_id']}"
        ws.cell(r, col["ticket_id"], it["ticket_id"])
        ws.cell(r, col["bedingung"], it["bedingung"])
        ws.cell(r, col["kb_abdeckung"], it["kb_abdeckung"])
        ws.cell(r, col["schwierigkeit"], it["schwierigkeit"])
        if llm:  # K-02: keine Belege, keine Fehlertypen bei LLM-only
            for k in ("K1_C2", "K2_C2", "K3_C2"):
                ws.cell(r, col[k], "kein Beleg").fill = grey
            for k in ("mangel", "F_primaer", "F_weitere"):
                ws.cell(r, col[k], NA).fill = grey
        for k in ("K1", "K2", "K3", "bemerkung"):
            ws.cell(r, col[k]).alignment = Alignment(wrap_text=True, vertical="top")

    def letter(name):
        return ws.cell(1, col[name]).column_letter

    last = len(items) + 1
    rag_rows = [r for r, it in enumerate(items, start=2) if it["bedingung"] == "rag"]

    def add(values, names, rows=None):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.error = "Bitte einen Wert aus der Liste wählen (siehe Blatt Codes)."
        ws.add_data_validation(dv)
        for name in names:
            if rows is None:
                dv.add(f"{letter(name)}2:{letter(name)}{last}")
            else:
                for r in rows:
                    dv.add(f"{letter(name)}{r}")

    add(C1, ["K1_C1", "K2_C1", "K3_C1"])
    add(C2, ["K1_C2", "K2_C2", "K3_C2"], rag_rows)
    add(JA_NEIN, ["C3"])
    add(JA_NEIN, ["mangel"], rag_rows)
    add(FEHLER, ["F_primaer"], rag_rows)
    widths = {"kodier_id": 10, "ticket_id": 9, "bedingung": 9, "kb_abdeckung": 11, "schwierigkeit": 11,
              "K1": 40, "K2": 40, "K3": 40, "F_weitere": 12, "bemerkung": 40}
    for name in cols:
        ws.column_dimensions[letter(name)].width = widths.get(name, 13)
    ws.freeze_panes = "F2"

    codes = wb.create_sheet("Codes")
    codes.append(["Spalte", "Werte", "Kurzregel (Details: LEITFADEN.md)"])
    for c in codes[1]:
        c.font = Font(bold=True)
    for row in [
        ("K1–K3", "Freitext", "bis zu drei Kernaussagen, knapp paraphrasiert; leer, wenn keine Kernaussage vorliegt"),
        ("Kx_C1", " / ".join(C1), "Ist die Kernaussage durch die gesamte KB gedeckt?"),
        ("Kx_C2", " / ".join(C2), "Stützt der im Satz angegebene Beleg die Aussage? LLM-only: immer „kein Beleg“"),
        ("C3", " / ".join(JA_NEIN), "Tritt eine Rückfrage, Enthaltung oder Weiterleitung an die Stelle einer Sachaussage zu einem Anliegen?"),
        ("mangel", " / ".join(JA_NEIN), "Nur RAG: ungedeckte Aussage, unbeantwortetes Anliegen (KB-gedeckt) oder Belegproblem?"),
        ("F_primaer", " / ".join(FEHLER), "Nur RAG und nur bei Mangel: erster zutreffender Typ in der Reihenfolge F1 → F2 → F3 → F4"),
        ("F_weitere", "z. B. F3;F4", "weitere zutreffende Typen, mit Semikolon getrennt"),
        ("bemerkung", "Freitext", "Begründung bei Unsicherheit oder Mischfall"),
    ]:
        codes.append(row)
    codes.column_dimensions["A"].width = 12
    codes.column_dimensions["B"].width = 34
    codes.column_dimensions["C"].width = 95
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--run", default="haupt-01")
    ap.add_argument("--out", type=Path, default=KODIER_DIR)
    args = ap.parse_args()
    out = args.out if args.out.is_absolute() else ROOT / args.out
    git_state = git_info()

    responses = ROOT / "rag" / "runs" / args.run / "responses.jsonl"
    with responses.open(encoding="utf-8") as f:
        entries = [json.loads(line) for line in f]
    with (ROOT / "tickets.jsonl").open(encoding="utf-8") as f:
        tickets = {t["ticket_id"]: t for t in map(json.loads, f)}
    kb = load_kb()

    keys = [(e["ticket_id"], e["bedingung"]) for e in entries]
    errors = []
    if len(set(keys)) != len(keys):
        errors.append("doppelte Einträge (Ticket, Bedingung) im Lauf")
    errors += [f"{t}: nicht in tickets.jsonl" for t, _ in keys if t not in tickets]
    errors += [f"{e['ticket_id']} rag: {len((e.get('retrieval') or {}).get('passagen', []))} Auszüge statt 3"
               for e in entries if e["bedingung"] == "rag" and len((e.get("retrieval") or {}).get("passagen", [])) != 3]
    if errors:
        sys.exit("FEHLER:\n" + "\n".join(errors))

    order = sorted(entries, key=lambda e: (e["ticket_id"], e["bedingung"]))
    random.Random(SEED).shuffle(order)
    items = [{"kodier_id": f"K{n:03d}", "ticket_id": e["ticket_id"], "bedingung": e["bedingung"],
              "kb_abdeckung": e["kb_abdeckung"], "schwierigkeit": e["schwierigkeit"], "entry": e}
             for n, e in enumerate(order, start=1)]

    out.mkdir(parents=True, exist_ok=True)
    (out / "kodierheft.html").write_text(build_html(items, tickets, kb, args.run), encoding="utf-8")
    build_xlsx(items, out / "kodiertabelle.xlsx")

    stichprobe = []
    if args.run.startswith("haupt"):
        rng = random.Random(f"{SEED}:zweitkodierung")
        for bed, n in N_ZWEIT.items():
            pool = [it for it in items if it["bedingung"] == bed]
            stichprobe += rng.sample(pool, n)
        stichprobe.sort(key=lambda it: it["kodier_id"])
        with (out / "stichprobe_zweitkodierung.csv").open("w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["kodier_id", "ticket_id", "bedingung"])
            w.writerows([it["kodier_id"], it["ticket_id"], it["bedingung"]] for it in stichprobe)

    meta = {"erzeugt": now_iso(), "skript": "kodierung/prepare_coding.py", "git": git_state, "seed": SEED,
            "lauf": args.run, "responses_sha256": file_hash(responses),
            "tickets_jsonl_sha256": file_hash(ROOT / "tickets.jsonl"),
            "anzahl": {"entwuerfe": len(items), "rag": sum(it["bedingung"] == "rag" for it in items),
                       "zweitkodierung": len(stichprobe)},
            "ausgaben": {n: file_hash(out / n) for n in ("kodierheft.html", "kodiertabelle.xlsx")}}
    (out / "kodier_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    rel = out.relative_to(ROOT)
    print(f"{len(items)} Entwürfe aus {args.run} ({meta['anzahl']['rag']} RAG) in zufälliger Reihenfolge (Seed {SEED})")
    print(f"Geschrieben: {rel}/kodierheft.html, {rel}/kodiertabelle.xlsx, {rel}/kodier_meta.json"
          + (f", {rel}/stichprobe_zweitkodierung.csv ({len(stichprobe)} Entwürfe)" if stichprobe else ""))


if __name__ == "__main__":
    main()
