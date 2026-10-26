"""Überträgt die Konsensentscheidungen (K-11) auf die Erstkodierung und erzeugt den endgültigen Kodierdatensatz.

Voraussetzungen:
- kodierung/kappa/abweichungen.xlsx aus kappa.py, Spalten „konsens“ und „begruendung“ für jede Zeile ausgefüllt
- kodierung/kodiertabelle.xlsx unverändert seit kappa.py (Prüfung über kappa_meta.json)

Aufruf (im Repo-Hauptverzeichnis):
    python kodierung/konsens.py

Ausgabe in kodierung/:
- kodiertabelle_konsens.xlsx   Erstkodierung mit den Konsenscodes für die 48 Entwürfe der Zweitkodierung;
                               Blatt „Konsens“ protokolliert jede Entscheidung, geänderte Zellen sind grün markiert
- kappa/konsens_meta.json      Herkunft: Hashes der Eingaben und der Ausgabe, Git-Stand

Grundlage aller Auswertungen von Strang 4 ist danach kodiertabelle_konsens.xlsx. Kappa bleibt auf den Codes vor dem
Konsens berechnet (kappa/bericht.md).
"""

import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from prepare_coding import KODIER_DIR, ROOT, git_info, now_iso  # noqa: E402
from pruefe_kodierung import C1, C2, FT, JA_NEIN, pruefe, s, split_f  # noqa: E402

ERST = KODIER_DIR / "kodiertabelle.xlsx"
ZWEIT = KODIER_DIR / "zweitkodierung" / "kodiertabelle_zweit.xlsx"
KAPPA_DIR = KODIER_DIR / "kappa"
ABW = KAPPA_DIR / "abweichungen.xlsx"
KAPPA_META = KAPPA_DIR / "kappa_meta.json"
OUT = KODIER_DIR / "kodiertabelle_konsens.xlsx"
GRUEN = PatternFill("solid", fgColor="E2EFDA")
LEER_F = {"", "keine", "-", "–"}


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def gueltig(code: str, wert: str) -> bool:
    if re.fullmatch(r"K[123]_C1", code):
        return wert in C1
    if re.fullmatch(r"K[123]_C2", code):
        return wert in C2
    if code in ("C3", "mangel"):
        return wert in JA_NEIN
    if code == "F_primaer":
        return wert in FT + ["keiner"]
    if code == "F_weitere":
        return wert == "" or all(x in FT for x in split_f(wert))
    return False


def main() -> None:
    meta = json.loads(KAPPA_META.read_text(encoding="utf-8"))
    fehler = []
    if file_hash(ERST) != meta["erstkodierung_sha256"]:
        fehler.append("kodiertabelle.xlsx hat sich seit kappa.py geändert; erst kappa.py erneut ausführen")
    if file_hash(ZWEIT) != meta["zweitkodierung_sha256"]:
        fehler.append("kodiertabelle_zweit.xlsx hat sich seit kappa.py geändert; erst kappa.py erneut ausführen")
    if fehler:
        sys.exit("FEHLER:\n" + "\n".join(fehler))

    zweit_ids = {s(d["kodier_id"]) for d in pruefe(ZWEIT)[0]}
    ws_abw = openpyxl.load_workbook(ABW)["Abweichungen"]
    rows = list(ws_abw.iter_rows(values_only=True))
    head = [s(h) for h in rows[0]]
    eintraege = [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]

    entscheidungen = {}
    for n, r in enumerate(eintraege, start=2):
        kid, code, kons, grund = s(r.get("kodier_id")), s(r.get("code")), s(r.get("konsens")), s(r.get("begruendung"))
        erst_wert, zweit_wert = s(r.get("erstkodierung")), s(r.get("zweitkodierung"))
        ort = f"Zeile {n} ({kid} {code})"
        aus_kappa = s(r.get("nr")) != ""
        if kid not in zweit_ids:
            fehler.append(f"{ort}: Entwurf gehört nicht zur Zweitkodierung")
            continue
        if not kons:
            fehler.append(f"{ort}: „konsens“ ist leer")
            continue
        if not grund:
            fehler.append(f"{ort}: „begruendung“ ist leer")
        if kons in ("erst", "zweit"):
            if not aus_kappa:
                fehler.append(f"{ort}: „erst“/„zweit“ nur in Zeilen aus kappa.py zulässig")
                continue
            kons = erst_wert if kons == "erst" else zweit_wert
        if code == "F_weitere":
            kons = "" if kons in LEER_F else ";".join(split_f(kons))
        if not gueltig(code, kons):
            fehler.append(f"{ort}: ungültiger Konsenscode {kons!r}")
            continue
        if (kid, code) in entscheidungen:
            fehler.append(f"{ort}: für {kid} {code} gibt es schon eine Zeile; dort den Konsens anpassen statt eine neue anzulegen")
            continue
        entscheidungen[(kid, code)] = {"konsens": kons, "begruendung": grund, "zweit": zweit_wert,
                                       "herkunft": "Abweichung" if aus_kappa else "Folgeänderung"}
    if fehler:
        sys.exit(f"FEHLER in {ABW.relative_to(ROOT)}:\n" + "\n".join(fehler))

    wb = openpyxl.load_workbook(ERST)
    ws = wb["Kodierung"]
    col = {s(c.value): c.column for c in ws[1]}
    zeile = {s(ws.cell(r, col["kodier_id"]).value): r for r in range(2, ws.max_row + 1)}
    protokoll = []
    for (kid, code), d in sorted(entscheidungen.items()):
        if code not in col:
            sys.exit(f"FEHLER: Spalte {code} fehlt in kodiertabelle.xlsx")
        if re.fullmatch(r"K[123]_C[12]", code) and not s(ws.cell(zeile[kid], col[code.split("_")[0]]).value):
            sys.exit(f"FEHLER: {kid} {code}: zu dieser Kernaussage gibt es keinen Text")
        cell = ws.cell(zeile[kid], col[code])
        alt = s(cell.value)
        neu = d["konsens"] or None
        if s(neu) != alt:
            cell.value = neu
            cell.fill = GRUEN
        protokoll.append([kid, code, d["herkunft"], alt, d["zweit"], s(neu), d["begruendung"]])

    if "Konsens" in wb.sheetnames:
        del wb["Konsens"]
    wk = wb.create_sheet("Konsens")
    wk.append(["kodier_id", "code", "herkunft", "erstkodierung", "zweitkodierung", "konsens", "begruendung"])
    for c in wk[1]:
        c.font = Font(bold=True)
    for p in protokoll:
        wk.append(p)
    for colname, wdt in zip("ABCDEFG", (9, 10, 14, 14, 14, 14, 70)):
        wk.column_dimensions[colname].width = wdt
    for row in wk.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT)

    _, fehler_neu, hinweise = pruefe(OUT)
    geaendert = sum(p[3] != p[5] for p in protokoll)
    print(f"{len(protokoll)} Entscheidungen übernommen, davon {geaendert} Änderungen gegenüber der Erstkodierung")
    if fehler_neu:
        print(f"\nFEHLER in {OUT.name} nach dem Konsens: {sum(len(v) for v in fehler_neu.values())}")
        for cat, items in fehler_neu.items():
            print(f"  {cat}: " + ", ".join(kid for kid, _ in items))
        print("\nFolgeänderungen in abweichungen.xlsx eintragen: Gibt es für den Code schon eine Zeile, dort den Konsens "
              "anpassen, sonst eine neue Zeile (kodier_id, code, konsens, begruendung). Danach konsens.py erneut ausführen.")
        sys.exit(1)
    n_h = sum(len(v) for v in hinweise.values())
    meta_out = {"erzeugt": now_iso(), "skript": "kodierung/konsens.py", "git": git_info(),
                "erstkodierung_sha256": file_hash(ERST), "zweitkodierung_sha256": file_hash(ZWEIT),
                "abweichungen_sha256": file_hash(ABW), "konsens_sha256": file_hash(OUT),
                "entscheidungen": len(protokoll), "aenderungen": geaendert}
    (KAPPA_DIR / "konsens_meta.json").write_text(json.dumps(meta_out, ensure_ascii=False, indent=2) + "\n",
                                                 encoding="utf-8")
    print(f"Prüfung: 0 Fehler, {n_h} Hinweise (python kodierung/pruefe_kodierung.py {OUT.relative_to(ROOT)})")
    print(f"Geschrieben: {OUT.relative_to(ROOT)}, {(KAPPA_DIR / 'konsens_meta.json').relative_to(ROOT)}")


if __name__ == "__main__":
    main()
