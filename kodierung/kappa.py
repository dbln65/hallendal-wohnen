"""Übereinstimmung von Erst- und Zweitkodierung (K-06, DEMI 2.5) und Liste der Abweichungen für den Konsens (K-11).

Voraussetzungen:
- kodierung/kodiertabelle.xlsx und kodierung/zweitkodierung/kodiertabelle_zweit.xlsx bestehen pruefe_kodierung.py
  ohne Fehler
- kodierung/zweitkodierung/zweit_meta.json aus zweitkodierung.py

Aufruf (im Repo-Hauptverzeichnis):
    python kodierung/kappa.py

Ausgabe in kodierung/kappa/:
- bericht.md          Cohens Kappa je Code mit prozentualer Übereinstimmung und Kreuztabellen
- kappa.csv           Kennwerte je Code
- abweichungen.xlsx   jede abweichende Codierung mit Spalten „konsens“ und „begruendung“ zum Ausfüllen
- kappa_meta.json     Herkunft: Hashes beider Kodiertabellen, Git-Stand

Kappa wird für die Codes vor dem Konsens berichtet. Analyseeinheit: C1 und C2 je Kernaussage, C3, Mangel und primärer
Fehlertyp je Entwurf (K-06). Mangel und Fehlertypen gibt es nur für RAG-Entwürfe (K-02).
"""

import csv
import hashlib
import json
import math
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from prepare_coding import KODIER_DIR, ROOT, git_info, now_iso  # noqa: E402
from pruefe_kodierung import pruefe, s, split_f  # noqa: E402

ERST = KODIER_DIR / "kodiertabelle.xlsx"
ZWEIT = KODIER_DIR / "zweitkodierung" / "kodiertabelle_zweit.xlsx"
ZWEIT_META = KODIER_DIR / "zweitkodierung" / "zweit_meta.json"
OUT = KODIER_DIR / "kappa"

ORDNUNG_C1 = ["gedeckt", "teilweise", "nicht gedeckt"]
KATEGORIEN = {
    "C1": ORDNUNG_C1,
    "C2": ["passend", "nicht passend", "kein Beleg"],
    "C3": ["ja", "nein"],
    "mangel": ["ja", "nein"],
    "F_primaer": ["keiner", "F1", "F2", "F3", "F4"],
}
BEZEICHNUNG = {"C1": "C1 Aussagendeckung", "C2": "C2 Quellenpassung", "C3": "C3 Rückfrage/Enthaltung",
               "mangel": "Mangel", "F_primaer": "Primärer Fehlertyp"}
EINHEIT = {"C1": "Kernaussagen (beide Bedingungen)", "C2": "Kernaussagen (RAG)", "C3": "Entwürfe (beide Bedingungen)",
           "mangel": "Entwürfe (RAG)", "F_primaer": "Entwürfe (RAG)"}
# Landis & Koch (1977), Originalbezeichnungen
STUFEN = [(0.0, "poor"), (0.20, "slight"), (0.40, "fair"), (0.60, "moderate"), (0.80, "substantial"),
          (1.0, "almost perfect")]


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def de(x, nd=2) -> str:
    return f"{x:,.{nd}f}".replace(",", "§").replace(".", ",").replace("§", ".").replace("-", "−")


def stufe(k: float) -> str:
    if k < 0:
        return "poor"
    for grenze, name in STUFEN[1:]:
        if k <= grenze + 1e-12:
            return name
    return "almost perfect"


def kappa(paare: list, kategorien: list, gewichtet: bool = False) -> dict:
    """Cohens Kappa für zwei Kodierende; optional linear gewichtet (nur für ordinale Kategorien sinnvoll).

    Standardfehler nach der Näherung von Cohen (1960): sqrt(po·(1−po) / (n·(1−pe)²)), nur ungewichtet.
    """
    n = len(paare)
    if n == 0:
        return {"n": 0}
    idx = {c: i for i, c in enumerate(kategorien)}
    m = len(kategorien)
    w = [[1 - abs(i - j) / (m - 1) if gewichtet else float(i == j) for j in range(m)] for i in range(m)]
    tab = [[0] * m for _ in range(m)]
    for a, b in paare:
        tab[idx[a]][idx[b]] += 1
    zeile = [sum(tab[i]) / n for i in range(m)]
    spalte = [sum(tab[i][j] for i in range(m)) / n for j in range(m)]
    po = sum(w[i][j] * tab[i][j] for i in range(m) for j in range(m)) / n
    pe = sum(w[i][j] * zeile[i] * spalte[j] for i in range(m) for j in range(m))
    roh = sum(tab[i][i] for i in range(m)) / n
    if abs(1 - pe) < 1e-12:
        return {"n": n, "po": po, "uebereinstimmung": roh, "pe": pe, "kappa": None, "tabelle": tab}
    k = (po - pe) / (1 - pe)
    res = {"n": n, "po": po, "uebereinstimmung": roh, "pe": pe, "kappa": k, "tabelle": tab}
    if not gewichtet:
        se = math.sqrt(po * (1 - po) / (n * (1 - pe) ** 2))
        res.update({"se": se, "ci_unten": max(-1.0, k - 1.96 * se), "ci_oben": min(1.0, k + 1.96 * se)})
    return res


def lade(path: Path, rolle: str) -> dict:
    data, fehler, _ = pruefe(path)
    if fehler:
        sys.exit(f"FEHLER: {path.relative_to(ROOT)} hat {sum(len(v) for v in fehler.values())} Fehler. "
                 f"Erst python kodierung/pruefe_kodierung.py {path.relative_to(ROOT)} ausführen; "
                 f"die {rolle} behebt sie.")
    return {s(d["kodier_id"]): d for d in data}


def main() -> None:
    erst, zweit = lade(ERST, "Erstkodierperson"), lade(ZWEIT, "Zweitkodierperson")
    meta_zweit = json.loads(ZWEIT_META.read_text(encoding="utf-8"))
    hinweise = []
    if file_hash(ERST) != meta_zweit.get("erstkodierung_sha256"):
        hinweise.append("kodiertabelle.xlsx hat sich seit zweitkodierung.py geändert; maßgeblich sind die "
                        "aktuellen Codes der Erstkodierung (Änderungen siehe Blatt „Korrekturen“).")

    fehler = []
    if len(zweit) != meta_zweit["anzahl"]["entwuerfe"]:
        fehler.append(f"Zweitkodierung enthält {len(zweit)} statt {meta_zweit['anzahl']['entwuerfe']} Entwürfe")
    for kid, z in zweit.items():
        e = erst.get(kid)
        if e is None:
            fehler.append(f"{kid}: fehlt in der Erstkodierung")
            continue
        for col in ("ticket_id", "bedingung"):
            if s(e[col]) != s(z[col]):
                fehler.append(f"{kid}: {col} weicht ab ({s(e[col])} / {s(z[col])})")
        for i in (1, 2, 3):
            if s(e[f"K{i}"]) != s(z[f"K{i}"]):
                fehler.append(f"{kid}: Kernaussage K{i} weicht von der Erstkodierung ab (vorgegeben, K-06)")
    if fehler:
        sys.exit("FEHLER:\n" + "\n".join(fehler))

    ids = sorted(zweit)
    paare = {code: [] for code in KATEGORIEN}
    abw = []

    def notiere(kid, code, kernaussage, a, b):
        if a != b:
            e = erst[kid]
            abw.append({"kodier_id": kid, "ticket_id": s(e["ticket_id"]), "bedingung": s(e["bedingung"]),
                        "kb_abdeckung": s(e["kb_abdeckung"]), "code": code, "kernaussage": kernaussage,
                        "erstkodierung": a, "zweitkodierung": b, "bemerkung_erst": s(e["bemerkung"]),
                        "bemerkung_zweit": s(zweit[kid]["bemerkung"])})

    for kid in ids:
        e, z = erst[kid], zweit[kid]
        rag = s(e["bedingung"]) == "rag"
        for i in (1, 2, 3):
            k = s(e[f"K{i}"])
            if not k:
                continue
            a, b = s(e[f"K{i}_C1"]), s(z[f"K{i}_C1"])
            paare["C1"].append((a, b))
            notiere(kid, f"K{i}_C1", k, a, b)
            if rag:
                a, b = s(e[f"K{i}_C2"]), s(z[f"K{i}_C2"])
                paare["C2"].append((a, b))
                notiere(kid, f"K{i}_C2", k, a, b)
        a, b = s(e["C3"]), s(z["C3"])
        paare["C3"].append((a, b))
        notiere(kid, "C3", "", a, b)
        if rag:
            for code in ("mangel", "F_primaer"):
                a, b = s(e[code]), s(z[code])
                paare[code].append((a, b))
                notiere(kid, code, "", a, b)
            a, b = ";".join(split_f(s(e["F_weitere"]))), ";".join(split_f(s(z["F_weitere"])))
            notiere(kid, "F_weitere", "", a, b)

    erg = {code: kappa(p, KATEGORIEN[code]) for code, p in paare.items()}
    erg_c1_w = kappa(paare["C1"], ORDNUNG_C1, gewichtet=True)

    OUT.mkdir(exist_ok=True)
    with (OUT / "kappa.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["code", "einheit", "n", "uebereinstimmung", "kappa", "ci_unten", "ci_oben", "stufe"])
        for code, r in erg.items():
            k = r.get("kappa")
            w.writerow([code, EINHEIT[code], r["n"], round(r["uebereinstimmung"], 4),
                        "" if k is None else round(k, 4), round(r["ci_unten"], 4) if k is not None else "",
                        round(r["ci_oben"], 4) if k is not None else "", stufe(k) if k is not None else ""])
        w.writerow(["C1_linear_gewichtet", EINHEIT["C1"], erg_c1_w["n"], round(erg_c1_w["po"], 4),
                    round(erg_c1_w["kappa"], 4) if erg_c1_w.get("kappa") is not None else "", "", "",
                    stufe(erg_c1_w["kappa"]) if erg_c1_w.get("kappa") is not None else ""])

    # Abweichungsliste für den Konsens
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Abweichungen"
    kopf = ["nr", "kodier_id", "ticket_id", "bedingung", "kb_abdeckung", "code", "kernaussage", "erstkodierung",
            "zweitkodierung", "bemerkung_erst", "bemerkung_zweit", "konsens", "begruendung"]
    ws.append(kopf)
    gelb = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True)
    for nr, a in enumerate(abw, start=1):
        ws.append([nr] + [a[k] for k in kopf[1:-2]] + [None, None])
        for col in (12, 13):
            ws.cell(ws.max_row, col).fill = gelb
    for col, wdt in zip("ABCDEFGHIJKLM", (5, 9, 9, 9, 11, 10, 60, 14, 14, 30, 30, 14, 50)):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    erlaubt = sorted({v for vals in KATEGORIEN.values() for v in vals} | {"erst", "zweit"})
    dv = DataValidation(type="list", formula1='"' + ",".join(erlaubt) + '"', allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    if abw:
        dv.add(f"L2:L{len(abw) + 1}")
    info = wb.create_sheet("Anleitung")
    for zeile in [
        "Konsens (K-11): Jede Zeile gemeinsam anhand des Kodierhefts (zweitkodierung/kodierheft.html), "
        "LEITFADEN.md und AUSLEGUNGEN.md besprechen.",
        "Spalte „konsens“: vereinbarter Code. Abkürzungen: „erst“ = Code der Erstkodierung, „zweit“ = Code der "
        "Zweitkodierung. Bei F_weitere mehrere Typen mit Semikolon (z. B. F3;F4), leer = keine.",
        "Spalte „begruendung“: kurzer Grund mit Verweis auf Leitfaden oder Auslegung (z. B. „A-08: vage, aber "
        "ableitbar“).",
        "Folgeänderungen: Ändert der Konsens einen Code so, dass ein anderer Code des Entwurfs angepasst werden "
        "muss (z. B. C1 → teilweise, dann Mangel = ja): Gibt es für diesen Code schon eine Zeile, dort den Konsens "
        "anpassen, sonst eine neue Zeile mit kodier_id, code, konsens und begruendung ergänzen. konsens.py meldet "
        "solche Fälle.",
        "Danach: python kodierung/konsens.py",
    ]:
        info.append([zeile])
    info.column_dimensions["A"].width = 140
    for row in info.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT / "abweichungen.xlsx")

    # Bericht
    L = ["# Übereinstimmung der Zweitkodierung", "",
         f"Erzeugt mit `kodierung/kappa.py` am {now_iso()[:10]}. {len(ids)} Entwürfe "
         f"({sum(s(erst[k]['bedingung']) == 'rag' for k in ids)} RAG), Codes vor dem Konsens (K-06, K-11).", ""]
    for h in hinweise:
        L += [f"Hinweis: {h}", ""]
    L += ["## Tabelle: Cohens Kappa je Code", "",
          "| Code | Einheit | n | Übereinstimmung | Kappa | 95-%-KI | Stufe |",
          "| --- | --- | --- | --- | --- | --- | --- |"]
    for code, r in erg.items():
        k = r.get("kappa")
        if k is None:
            L.append(f"| {BEZEICHNUNG[code]} | {EINHEIT[code]} | {r['n']} | {de(r['uebereinstimmung'] * 100, 1)} % "
                     "| nicht definiert¹ | – | – |")
        else:
            L.append(f"| {BEZEICHNUNG[code]} | {EINHEIT[code]} | {r['n']} | {de(r['uebereinstimmung'] * 100, 1)} % | "
                     f"{de(k)} | {de(r['ci_unten'])} bis {de(r['ci_oben'])} | {stufe(k)} |")
    if erg_c1_w.get("kappa") is not None:
        L.append(f"| C1, linear gewichtet (ergänzend) | {EINHEIT['C1']} | {erg_c1_w['n']} | – | "
                 f"{de(erg_c1_w['kappa'])} | – | {stufe(erg_c1_w['kappa'])} |")
    L += ["",
          "Stufen nach Landis und Koch (1977). Konfidenzintervall nach der Näherung von Cohen (1960), auf −1 bis 1 begrenzt; "
          "bei kleinen Fallzahlen nur grob. Bei sehr ungleich verteilten Kategorien kann Kappa trotz hoher Übereinstimmung niedrig "
          "ausfallen; deshalb wird die prozentuale Übereinstimmung mit berichtet. Die gewichtete Variante wertet "
          "„teilweise“ gegenüber „gedeckt“ oder „nicht gedeckt“ als halbe Übereinstimmung.",
          "",
          "¹ Kappa ist nicht definiert, wenn beide Kodierenden durchgehend dieselbe Kategorie vergeben.",
          "", "## Kreuztabellen", "", "Zeilen: Erstkodierung, Spalten: Zweitkodierung.", ""]
    for code, r in erg.items():
        kat = KATEGORIEN[code]
        L += [f"**{BEZEICHNUNG[code]}** (n = {r['n']})", "",
              "| Erst \\ Zweit | " + " | ".join(kat) + " |", "| --- |" + " --- |" * len(kat)]
        for i, a in enumerate(kat):
            L.append(f"| {a} | " + " | ".join(str(r["tabelle"][i][j]) for j in range(len(kat))) + " |")
        L.append("")
    zaehl = Counter(a["code"].split("_", 1)[-1] if a["code"].startswith("K") else a["code"] for a in abw)
    L += ["## Abweichungen für den Konsens", "",
          f"{len(abw)} abweichende Codierungen in {len({a['kodier_id'] for a in abw})} Entwürfen "
          f"(davon F_weitere: {zaehl.get('F_weitere', 0)}; nicht Teil der Kappa-Berechnung). "
          "Liste: `abweichungen.xlsx`.", ""]
    (OUT / "bericht.md").write_text("\n".join(L), encoding="utf-8")

    meta = {"erzeugt": now_iso(), "skript": "kodierung/kappa.py", "git": git_info(),
            "erstkodierung_sha256": file_hash(ERST), "zweitkodierung_sha256": file_hash(ZWEIT),
            "zweit_meta_sha256": file_hash(ZWEIT_META), "entwuerfe": len(ids), "abweichungen": len(abw),
            "abweichungen_sha256": file_hash(OUT / "abweichungen.xlsx"),
            "kappa": {c: (None if r.get("kappa") is None else round(r["kappa"], 4)) for c, r in erg.items()}}
    (OUT / "kappa_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print("\n".join(L[:4 + 2 * len(hinweise)]))
    for code, r in erg.items():
        k = r.get("kappa")
        print(f"  {BEZEICHNUNG[code]:<24} n = {r['n']:>3}  Übereinstimmung {de(r['uebereinstimmung'] * 100, 1):>5} %  "
              f"Kappa {'nicht definiert' if k is None else de(k)}")
    print(f"\n{len(abw)} Abweichungen → {(OUT / 'abweichungen.xlsx').relative_to(ROOT)}")
    print(f"Geschrieben: {OUT.relative_to(ROOT)}/ (bericht.md, kappa.csv, abweichungen.xlsx, kappa_meta.json)")


if __name__ == "__main__":
    main()
